# orders/reports/sheets/inconsistencies_sheet.py
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from .base_sheet import BaseSheet
from ..styles.theme import FILLS, FONTS, BORDERS, ALIGNMENTS, COLORS  

class InconsistenciesSheet(BaseSheet):
    """Лист с детализацией несоответствий"""
    
    def __init__(self, workbook, sheet_number):
        super().__init__(workbook, sheet_number)  
    
    def build(self, inconsistencies):
        """Построить лист с несоответствиями"""
        
        # Используем стандартный метод для заголовка
        self.add_header(
            title="⚠️ НЕСООТВЕТСТВИЯ В СТАТУСАХ ЗАКАЗОВ",
            subtitle="Отмененные заказы с активными статусами",
            note="Требуется проверка и исправление статусов"
        )
        
        # Добавляем кнопку возврата в оглавление
        self.add_toc_button()
        
        row = self.current_row
        
        # -------------------------------------------------
        # БЛОК 1: НА СОГЛАСОВАНИИ (ОТМЕНЕННЫЕ)
        # -------------------------------------------------
        if inconsistencies['agreement_cancelled']:
            # Заголовок блока
            self.add_section_title(
                "1. ЗАКАЗЫ СО СТАТУСОМ 'НА СОГЛАСОВАНИИ' (ОТМЕНЕННЫЕ)",
                col_start=1,
                col_end=5
            )
            
            # Заголовки таблицы
            headers = ["№", "НОМЕР ЗАКАЗА", "НАЗВАНИЕ", "ПРИЧИНА ОТМЕНЫ"]
            self.add_table_header(headers, start_col=1)
            
            # Данные
            for idx, order in enumerate(inconsistencies['agreement_cancelled'], 1):
                values = [
                    idx,
                    order.get('number', order.get('id', '')),
                    order.get('fullname', '')[:50],
                    order.get('cancellation_reason', 'Не указана')
                ]
                self.add_data_row(values, start_col=1)
            
            self.add_blank_row()
        
        # -------------------------------------------------
        # БЛОК 2: К ВЫПОЛНЕНИЮ / В РЕЗЕРВЕ (ОТМЕНЕННЫЕ)
        # -------------------------------------------------
        if inconsistencies['execution_cancelled']:
            # Заголовок блока
            self.add_section_title(
                "2. ЗАКАЗЫ СО СТАТУСОМ 'К ВЫПОЛНЕНИЮ / В РЕЗЕРВЕ' (ОТМЕНЕННЫЕ)",
                col_start=1,
                col_end=5
            )
            
            # Заголовки таблицы
            headers = ["№", "НОМЕР ЗАКАЗА", "НАЗВАНИЕ", "ПРИЧИНА ОТМЕНЫ"]
            self.add_table_header(headers, start_col=1)
            
            # Данные
            for idx, order in enumerate(inconsistencies['execution_cancelled'], 1):
                values = [
                    idx,
                    order.get('number', order.get('id', '')),
                    order.get('fullname', '')[:50],
                    order.get('cancellation_reason', 'Не указана')
                ]
                self.add_data_row(values, start_col=1)
        
        # Настройка ширины колонок
        self.set_column_widths({
            'A': 5,   # №
            'B': 18,  # НОМЕР ЗАКАЗА
            'C': 45,  # НАЗВАНИЕ
            'D': 35   # ПРИЧИНА ОТМЕНЫ
        })
        
        # Замораживаем шапку и скрываем сетку
        self.freeze_and_hide_grid("A6")