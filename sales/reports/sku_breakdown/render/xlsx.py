# sales/reports/sku_breakdown/render/xlsx.py
"""
Рендер отчёта "Разбивка по номенклатуре и штрихкодам" в .xlsx.

Три листа:
  - "Детализация"       — плоская таблица, всё как есть, сортировка по сумме;
  - "По категориям"     — та же детализация, сгруппированная по категориям,
                           с подытогом по каждой категории;
  - "По производителям" — то же самое, но сгруппировано по производителям.

Палитра — тот же бренд-язык COSMORELAX, что в newspaper/sales_digest
(sales.reports.newspaper.config), просто перенесённый на Excel вместо
HTML/CSS, чтобы отчёт не выглядел инородным среди остальных.
"""

from __future__ import annotations

from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sales.reports.newspaper.config import (
    COLOR_ACCENT,
    COLOR_ACCENT_SOFT,
    COLOR_BRAND_DEEP,
    COLOR_INK,
    COLOR_INK_SOFT,
    COLOR_PAPER,
    COLOR_RULE,
    COMPANY_BRAND_NAME,
)


def _hex(c: str) -> str:
    """openpyxl хочет ARGB/RGB без '#'."""
    return c.lstrip("#").upper()


FILL_HEADER = PatternFill("solid", fgColor=_hex(COLOR_BRAND_DEEP))
FILL_GROUP = PatternFill("solid", fgColor=_hex(COLOR_ACCENT_SOFT))
FILL_TOTAL = PatternFill("solid", fgColor=_hex(COLOR_PAPER))
FILL_GRAND_TOTAL = PatternFill("solid", fgColor=_hex(COLOR_ACCENT))

FONT_TITLE = Font(name="Calibri", size=15, bold=True, color=_hex(COLOR_BRAND_DEEP))
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color=_hex(COLOR_INK_SOFT))
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_GROUP = Font(name="Calibri", size=10.5, bold=True, color=_hex(COLOR_INK))
FONT_TOTAL = Font(name="Calibri", size=10, bold=True, color=_hex(COLOR_INK))
FONT_GRAND_TOTAL = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_CELL = Font(name="Calibri", size=10, color=_hex(COLOR_INK))

THIN = Side(style="thin", color=_hex(COLOR_RULE))
BORDER_ROW = Border(bottom=THIN)

MONEY_FMT = "#,##0.00"
QTY_FMT = "#,##0.###"

# (заголовок, ключ в DataFrame, ширина колонки, числовой формат)
METRIC_COLUMNS = [
    ("Продажи, сумма", "sales_amount", 15, MONEY_FMT),
    ("Продажи, кол-во", "sales_qty", 13, QTY_FMT),
    ("Возвраты, сумма", "returns_amount", 15, MONEY_FMT),
    ("Возвраты, кол-во", "returns_qty", 13, QTY_FMT),
    ("Итого, сумма", "net_amount", 15, MONEY_FMT),
    ("Итого, кол-во", "net_qty", 13, QTY_FMT),
]

DETAIL_COLUMNS = [
    ("Категория", "cat_name", 24, None),
    ("Подкатегория", "subcat_name", 22, None),
    ("Производитель", "manufacturer_name", 22, None),
    ("Артикул", "article", 14, None),
    ("Номенклатура", "fullname", 46, None),
    ("Штрихкод", "barcode", 16, None),
] + METRIC_COLUMNS

_REPORT_TYPE_LABEL = {"daily": "день", "weekly": "неделя", "monthly": "месяц"}


def _period_title(meta: dict) -> str:
    kind = _REPORT_TYPE_LABEL.get(meta["report_type"], meta["report_type"])
    s, e = meta["start"], meta["end"]
    if s == e:
        rng = s.strftime("%d.%m.%Y")
    else:
        rng = f'{s.strftime("%d.%m.%Y")} — {e.strftime("%d.%m.%Y")}'
    return f"Период: {kind}, {rng}"


def _write_title_block(ws: Worksheet, n_cols: int, subtitle: str, meta: dict) -> int:
    """Пишет строку заголовка + подзаголовок с периодом, возвращает номер
    строки, в которую нужно писать шапку таблицы (1-based)."""
    n_cols = max(n_cols, 1)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(
        row=1, column=1,
        value=f"{COMPANY_BRAND_NAME} — Разбивка по номенклатуре и штрихкодам",
    )
    title_cell.font = FONT_TITLE
    ws.row_dimensions[1].height = 22

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    subtitle_text = (
        f"{subtitle}   ·   {_period_title(meta)}   ·   "
        f"сформировано {meta['generated_at'].strftime('%d.%m.%Y %H:%M')}"
    )
    sub_cell = ws.cell(row=2, column=1, value=subtitle_text)
    sub_cell.font = FONT_SUBTITLE

    return 4


def _write_header(ws: Worksheet, row: int, columns) -> None:
    for idx, (title, _key, width, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=idx, value=title)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[row].height = 28


def _write_data_row(ws: Worksheet, row: int, columns, record: dict) -> None:
    for idx, (_title, key, _width, fmt) in enumerate(columns, start=1):
        value = record.get(key, "")
        if key == "barcode" and value is not None:
            value = str(value)
        cell = ws.cell(row=row, column=idx, value=value)
        cell.font = FONT_CELL
        cell.border = BORDER_ROW
        if key == "barcode":
            cell.number_format = "@"
        elif fmt:
            cell.number_format = fmt
        if key == "fullname":
            cell.alignment = Alignment(vertical="center")


def _sum_metrics(df: pd.DataFrame) -> dict:
    return {key: float(df[key].sum()) for _t, key, _w, _f in METRIC_COLUMNS}


def _write_total_row(ws: Worksheet, row: int, columns, label: str, totals: dict, *, grand: bool = False) -> None:
    fill = FILL_GRAND_TOTAL if grand else FILL_TOTAL
    font = FONT_GRAND_TOTAL if grand else FONT_TOTAL

    n_label_cols = max(len(columns) - len(METRIC_COLUMNS), 1)
    if n_label_cols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_label_cols)
    for col_idx in range(1, n_label_cols + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = fill
        cell.font = font
    ws.cell(row=row, column=1, value=label)

    for idx, (_title, key, _width, fmt) in enumerate(columns, start=1):
        if key not in totals:
            continue
        cell = ws.cell(row=row, column=idx, value=totals[key])
        cell.font = font
        cell.fill = fill
        if fmt:
            cell.number_format = fmt
    ws.row_dimensions[row].height = 20


def _build_detail_sheet(wb: Workbook, df: pd.DataFrame, meta: dict) -> None:
    ws = wb.active
    ws.title = "Детализация"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _hex(COLOR_BRAND_DEEP)

    header_row = _write_title_block(
        ws, len(DETAIL_COLUMNS), "Полная детализация по номенклатуре и штрихкодам", meta,
    )
    _write_header(ws, header_row, DETAIL_COLUMNS)

    data = df.sort_values("net_amount", ascending=False)
    row = header_row + 1
    for _, rec in data.iterrows():
        _write_data_row(ws, row, DETAIL_COLUMNS, rec.to_dict())
        row += 1

    _write_total_row(ws, row, DETAIL_COLUMNS, "ИТОГО ЗА ПЕРИОД", _sum_metrics(df), grand=True)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(DETAIL_COLUMNS))}{row - 1}"


def _build_grouped_sheet(
    wb: Workbook,
    df: pd.DataFrame,
    meta: dict,
    *,
    group_col: str,
    sheet_title: str,
    subtitle: str,
) -> None:
    columns = [c for c in DETAIL_COLUMNS if c[1] != group_col]
    ws = wb.create_sheet(sheet_title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _hex(COLOR_ACCENT)

    header_row = _write_title_block(ws, len(columns), subtitle, meta)

    group_totals = (
        df.groupby(group_col, dropna=False)["net_amount"]
        .sum()
        .sort_values(ascending=False)
    )

    row = header_row
    _write_header(ws, row, columns)
    row += 1

    for group_value in group_totals.index.tolist():
        group_df = df[df[group_col] == group_value].sort_values("net_amount", ascending=False)

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(columns))
        gcell = ws.cell(row=row, column=1, value=f"{group_value}  ({len(group_df)} поз.)")
        gcell.font = FONT_GROUP
        gcell.fill = FILL_GROUP
        for col_idx in range(1, len(columns) + 1):
            ws.cell(row=row, column=col_idx).fill = FILL_GROUP
        ws.row_dimensions[row].height = 20
        row += 1

        for _, rec in group_df.iterrows():
            _write_data_row(ws, row, columns, rec.to_dict())
            row += 1

        _write_total_row(ws, row, columns, f"Итого: {group_value}", _sum_metrics(group_df))
        row += 1

    _write_total_row(ws, row, columns, "ОБЩИЙ ИТОГ", _sum_metrics(df), grand=True)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def _build_empty_workbook(meta: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Детализация"
    ws.cell(row=1, column=1, value=f"{COMPANY_BRAND_NAME} — Разбивка по номенклатуре и штрихкодам").font = FONT_TITLE
    ws.cell(row=2, column=1, value=f"{_period_title(meta)} — за этот период продаж нет").font = FONT_SUBTITLE
    ws.column_dimensions["A"].width = 60
    return wb


def build_workbook(df: pd.DataFrame, *, report_type: str, start: date, end: date) -> Workbook:
    """Собирает готовую книгу Excel по данным из data.get_sku_barcode_breakdown."""
    meta = {
        "report_type": report_type,
        "start": start,
        "end": end,
        "generated_at": datetime.now(),
    }

    if df.empty:
        return _build_empty_workbook(meta)

    wb = Workbook()
    _build_detail_sheet(wb, df, meta)
    _build_grouped_sheet(
        wb, df, meta,
        group_col="cat_name", sheet_title="По категориям",
        subtitle="Сгруппировано по категориям",
    )
    _build_grouped_sheet(
        wb, df, meta,
        group_col="manufacturer_name", sheet_title="По производителям",
        subtitle="Сгруппировано по производителям",
    )
    wb.active = 0
    return wb
