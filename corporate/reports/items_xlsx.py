# corporate/reports/items_xlsx.py
"""Выгрузка номенклатуры из админки в Excel."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from sales.reports.xl_brand import (
    FMT_DATE,
    FMT_QTY,
    FMT_SHARE,
    build_toc,
    finalize,
    set_number_format,
    style_table_sheet,
)

from corporate.models import CatTree

NO_CAT = "Без категории"
NO_SUB = "Без подкатегории"

ITEM_COLUMNS = [
    ("root", "Группа", 18),
    ("cat", "Категория", 22),
    ("subcat", "Подкатегория", 22),
    ("fullname", "Номенклатура (1С)", 44),
    ("name", "Название на сайте", 36),
    ("article", "Артикул", 16),
    ("brend", "Бренд", 16),
    ("manufacturer", "Производитель", 20),
    ("collections", "Коллекции", 20),
    ("barcodes", "Штрихкоды", 22),
    ("onec_cat", "Группа 1С", 20),
    ("onec_subcat", "Вид 1С", 20),
    ("im_id", "ID на сайте", 12),
    ("init_date", "Дата появления", 13),
    ("id", "item_id", 10),
]


def _items_frame(queryset) -> pd.DataFrame:
    qs = (queryset.select_related("cat", "subcat", "manufacturer", "brend")
          .prefetch_related("barcode", "collection"))
    tree_ids = {it.cat.tree_id for it in qs if it.cat_id}
    roots = {r.tree_id: r.name for r in CatTree.objects.filter(tree_id__in=tree_ids, level=0)}

    rows = []
    for it in qs:
        rows.append({
            "root": roots.get(it.cat.tree_id, "") if it.cat_id else "",
            "cat": it.cat.name if it.cat_id else NO_CAT,
            "subcat": it.subcat.name if it.subcat_id else "",
            "fullname": it.fullname or "",
            "name": it.name or "",
            "article": it.article or "",
            "brend": it.brend.name if it.brend_id else "",
            "manufacturer": it.manufacturer.name if it.manufacturer_id else "",
            "collections": ", ".join(c.name for c in it.collection.all()),
            "barcodes": ", ".join(b.barcode for b in it.barcode.all()),
            "onec_cat": it.onec_cat or "",
            "onec_subcat": it.onec_subcat or "",
            "im_id": it.im_id or "",
            "init_date": it.init_date,
            "id": it.id,
        })
    df = pd.DataFrame(rows, columns=[c for c, _, _ in ITEM_COLUMNS])
    return df.sort_values(["root", "cat", "subcat", "fullname"], kind="stable").reset_index(drop=True)


def _write_df(ws, df: pd.DataFrame) -> None:
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)


def _widths(ws, widths) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_items_excel_bytes(queryset) -> bytes:
    df = _items_frame(queryset)
    total = len(df)
    no_cat = int((df["cat"] == NO_CAT).sum())
    no_sub = int(((df["cat"] != NO_CAT) & (df["subcat"] == "")).sum())
    no_manu = int((df["manufacturer"] == "").sum())

    wb = Workbook()
    wb.remove(wb.active)

    # Номенклатура
    ws = wb.create_sheet("Номенклатура")
    out = df.rename(columns={c: t for c, t, _ in ITEM_COLUMNS})
    _write_df(ws, out)
    hdr = style_table_sheet(
        ws, title="НОМЕНКЛАТУРА",
        subtitle=f"{total} позиций · сортировка: группа → категория → подкатегория",
        freeze_col=5, key_headers=("Номенклатура (1С)",),
    )
    set_number_format(ws, hdr, ["Дата появления"], FMT_DATE)
    _widths(ws, [w for _, _, w in ITEM_COLUMNS])

    # По категориям
    cat = df.assign(subcat=df["subcat"].replace("", NO_SUB))
    summary = (cat.groupby(["root", "cat", "subcat"], dropna=False)
               .agg(items=("id", "count"),
                    no_manu=("manufacturer", lambda s: int((s == "").sum())))
               .reset_index())
    summary["share"] = summary["items"] / total if total else 0.0
    summary = summary.rename(columns={
        "root": "Группа", "cat": "Категория", "subcat": "Подкатегория",
        "items": "Номенклатур", "no_manu": "Без производителя", "share": "Доля",
    })
    ws2 = wb.create_sheet("По категориям")
    _write_df(ws2, summary)
    hdr2 = style_table_sheet(ws2, title="НОМЕНКЛАТУРА ПО КАТЕГОРИЯМ",
                             subtitle="Количество позиций по группам, категориям и подкатегориям",
                             freeze_col=4, key_headers=("Номенклатур",))
    set_number_format(ws2, hdr2, ["Номенклатур", "Без производителя"], FMT_QTY)
    set_number_format(ws2, hdr2, ["Доля"], FMT_SHARE)
    _widths(ws2, [20, 26, 26, 14, 16, 10])

    # Требуют разметки
    todo = df[(df["cat"] == NO_CAT) | (df["subcat"] == "") | (df["manufacturer"] == "")].copy()
    if not todo.empty:
        todo.insert(0, "issue", [
            ", ".join(p for p, bad in (
                ("нет категории", r.cat == NO_CAT),
                ("нет подкатегории", r.cat != NO_CAT and r.subcat == ""),
                ("нет производителя", r.manufacturer == ""),
            ) if bad)
            for r in todo.itertuples()
        ])
        cols = [("issue", "Что заполнить", 30)] + [c for c in ITEM_COLUMNS
                                                   if c[0] in ("root", "cat", "subcat", "fullname", "article",
                                                               "manufacturer", "onec_cat", "onec_subcat", "id")]
        ws3 = wb.create_sheet("Требуют разметки")
        _write_df(ws3, todo[[c for c, _, _ in cols]].rename(columns={c: t for c, t, _ in cols}))
        style_table_sheet(ws3, title="ТРЕБУЮТ РАЗМЕТКИ",
                          subtitle="Позиции без категории, подкатегории или производителя",
                          freeze_col=6, key_headers=("Что заполнить",))
        _widths(ws3, [w for _, _, w in cols])

    build_toc(
        wb,
        title="Номенклатура",
        subtitle="Справочник товаров · категории, производители, штрихкоды",
        params=f"Выбрано позиций: {total} · выгружено {datetime.now():%d.%m.%Y %H:%M}",
        sheets=[
            ("Номенклатура", "Полный список выбранных позиций со всеми реквизитами"),
            ("По категориям", "Сколько позиций в каждой группе, категории и подкатегории"),
            ("Требуют разметки", "Позиции без категории, подкатегории или производителя"),
        ],
        cards=[
            ("НОМЕНКЛАТУР", total, FMT_QTY, "выбрано в админке"),
            ("БЕЗ КАТЕГОРИИ", no_cat, FMT_QTY, "нужно назначить"),
            ("БЕЗ ПОДКАТЕГОРИИ", no_sub, FMT_QTY, "с категорией, но без подкатегории"),
            ("БЕЗ ПРОИЗВОДИТЕЛЯ", no_manu, FMT_QTY, "нет в карточке"),
        ],
    )
    finalize(wb, ["Оглавление", "Номенклатура", "По категориям", "Требуют разметки"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
